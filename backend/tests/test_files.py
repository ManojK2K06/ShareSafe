import io,json,os,tempfile,unittest,zipfile
from pathlib import Path
from PIL import Image
from docx import Document
from sharesafe_engine.files import *
from sharesafe_engine.core import analyze,sanitize

class FileTests(unittest.TestCase):
    def test_utf8_import(self):self.assertEqual(ingest('example.txt','Hello 🙂'.encode())['text'],'Hello 🙂')
    def test_wrong_binary(self):
        with self.assertRaises(ValueError):ingest('x.txt',b'abc\x00def')
    def test_unsupported_type(self):
        with self.assertRaises(ValueError):ingest('x.exe',b'abc')
    def test_empty_upload(self):
        with self.assertRaises(ValueError):ingest('x.txt',b'')
    def test_bad_pdf(self):
        with self.assertRaises(ValueError):ingest('x.pdf',b'not pdf')
    def test_docx_split_runs_and_tables(self):
        d=Document();p=d.add_paragraph();p.add_run('alice@');p.add_run('example.com');d.add_table(rows=1,cols=1).cell(0,0).text='Account 12345678'
        b=io.BytesIO();d.save(b);asset=ingest('input.docx',b.getvalue());clean,_=sanitize(analyze(asset['text']))
        self.assertNotIn('alice@example.com',clean);self.assertNotIn('12345678',clean)
    def test_export_docx_no_hidden_source(self):
        data,_,_=export_text('[EMAIL_1]','docx')
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            self.assertFalse(any(b'alice@example.com' in z.read(n) for n in z.namelist()))
        self.assertEqual(Document(io.BytesIO(data)).paragraphs[0].text,'[EMAIL_1]')
    def test_pdf_import_and_clean_export(self):
        import pymupdf
        d=pymupdf.open();p=d.new_page();p.insert_text((50,50),'Email alice@example.com. Account 12345678.');raw=d.tobytes();d.close()
        a=ingest('test.pdf',raw);clean,_=sanitize(analyze(a['text']));data,_,_=export_text(clean,'pdf')
        with pymupdf.open(stream=data,filetype='pdf') as d:
            text=''.join(p.get_text() for p in d);self.assertNotIn('alice@example.com',text);self.assertNotIn('12345678',text)
    def test_pdf_unicode_escaped_not_lost(self):
        import pymupdf
        data,_,_=export_text('🙂 public','pdf')
        with pymupdf.open(stream=data,filetype='pdf') as d:self.assertIn('\\U0001f642',d[0].get_text())
    def test_image_pixels_permanently_redacted(self):
        im=Image.new('RGB',(100,100),'white');data=redact_image(im,[{'x':10,'y':10,'width':20,'height':20}]);clean=Image.open(io.BytesIO(data))
        self.assertEqual(clean.getpixel((15,15)),(0,0,0));self.assertEqual(clean.getpixel((50,50)),(255,255,255))
    def test_invalid_rectangle(self):
        with self.assertRaises(ValueError):redact_image(Image.new('RGB',(50,50)),[{'x':0,'y':0,'width':float('nan'),'height':2}])
    def test_png_metadata_removed(self):
        im=Image.new('RGB',(10,10));im.info['comment']='alice@example.com';data=redact_image(im,[])
        self.assertNotIn(b'alice@example.com',data)
    def test_exif_and_color_profile_metadata_not_carried_into_png(self):
        im=Image.new('RGB',(10,10));exif=im.getexif();exif[270]='alice@example.org';im.info['exif']=exif.tobytes();im.info['icc_profile']=b'fixture-private-profile'
        data=redact_image(im,[]);output=Image.open(io.BytesIO(data));self.assertNotIn('exif',output.info);self.assertNotIn('icc_profile',output.info);self.assertNotIn(b'alice@example.org',data)
    def test_scanned_pdf_no_exportable_text(self):
        import pymupdf
        d=pymupdf.open();d.new_page();data=d.tobytes();d.close()
        with self.assertRaises(ValueError):ingest('scan.pdf',data)

if __name__=='__main__':unittest.main()
